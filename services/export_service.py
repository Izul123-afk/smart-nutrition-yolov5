from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from io import BytesIO, StringIO
import csv
import html
import json
import re
import zipfile


@dataclass
class ExportArtifact:
    content: BytesIO
    filename: str
    mimetype: str


class ExportService:
    SINGLE_CSV_FIELDNAMES = [
        "nama",
        "makanan",
        "confidence_percent",
        "berat_gram",
        "protein_g",
        "lemak_g",
        "karbohidrat_g",
        "energi_kkal",
    ]

    def __init__(self, database_service) -> None:
        self.database_service = database_service

    @staticmethod
    def _filename(extension: str) -> str:
        timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        return f"food-recognition-{timestamp}.{extension}"

    @staticmethod
    def _flat_scan(record: dict) -> dict:
        timing = record.get("timings") or {}
        nutrition = record.get("nutrition") or {}
        per_100g = nutrition.get("per_100g") or {}
        totals = nutrition.get("totals") or {}
        return {
            "scan_id": record.get("scan_id"),
            "timestamp": record.get("timestamp"),
            "source": record.get("source"),
            "detected": record.get("detected"),
            "food_name": record.get("food_name"),
            "confidence_percent": record.get("confidence_percent"),
            "confidence_ratio": record.get("confidence_ratio"),
            "weight_grams": record.get("weight_grams"),
            "object_count": record.get("object_count"),
            "model_name": record.get("model_name"),
            "model_input": record.get("model_input"),
            "capture_ms": timing.get("capture_ms"),
            "preprocess_ms": timing.get("preprocess_ms"),
            "inference_ms": timing.get("inference_ms"),
            "postprocess_ms": timing.get("postprocess_ms"),
            "weight_ms": timing.get("weight_ms"),
            "complete_ms": timing.get("complete_ms"),
            "nutrition_available": nutrition.get("available"),
            "nutrition_data_version": nutrition.get("data_version"),
            "energy_kcal_per_100g": per_100g.get("energy_kcal"),
            "protein_g_per_100g": per_100g.get("protein_g"),
            "fat_g_per_100g": per_100g.get("fat_g"),
            "carbohydrate_g_per_100g": per_100g.get("carbohydrate_g"),
            "energy_kcal_total": totals.get("energy_kcal"),
            "protein_g_total": totals.get("protein_g"),
            "fat_g_total": totals.get("fat_g"),
            "carbohydrate_g_total": totals.get("carbohydrate_g"),
            "original_image_path": record.get("original_image_path"),
            "annotated_image_path": record.get("annotated_image_path"),
        }

    @staticmethod
    def _order_scans(
        scans: list[dict], scan_ids: list[str] | tuple[str, ...] | None
    ) -> list[dict]:
        """Urutkan hasil sesuai urutan pilihan user bila `scan_ids` diisi."""
        if not scan_ids:
            return scans
        order = {str(value): index for index, value in enumerate(scan_ids)}
        return sorted(
            scans, key=lambda item: order.get(str(item.get("scan_id")), 10**9)
        )

    def csv(self, since=None, scan_id=None, scan_ids=None) -> ExportArtifact:
        dataset = self.database_service.export_dataset(
            since=since, scan_id=scan_id, scan_ids=scan_ids
        )
        detections_by_scan: dict[str, list[dict]] = {}
        for detection in dataset["detections"]:
            detections_by_scan.setdefault(detection["scan_id"], []).append(
                detection
            )

        records = []
        for scan in dataset["scans"]:
            row = self._flat_scan(scan)
            row["detections_json"] = json.dumps(
                detections_by_scan.get(scan["scan_id"], []),
                ensure_ascii=False,
                separators=(",", ":"),
            )
            records.append(row)

        output = StringIO(newline="")
        fieldnames = list(records[0]) if records else list(
            self._flat_scan({}).keys()
        ) + ["detections_json"]
        writer = csv.DictWriter(output, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(records)
        content = BytesIO(("\ufeff" + output.getvalue()).encode("utf-8"))
        return ExportArtifact(
            content=content,
            filename=self._filename("csv"),
            mimetype="text/csv; charset=utf-8",
        )

    def xlsx(self, since=None, scan_id=None, scan_ids=None) -> ExportArtifact:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        dataset = self.database_service.export_dataset(
            since=since, scan_id=scan_id, scan_ids=scan_ids
        )
        workbook = Workbook()
        scans_sheet = workbook.active
        scans_sheet.title = "Scans"
        scan_rows = [self._flat_scan(item) for item in dataset["scans"]]
        self._fill_sheet(scans_sheet, scan_rows, Font, PatternFill)
        detections_sheet = workbook.create_sheet("Detections")
        self._fill_sheet(
            detections_sheet,
            dataset["detections"],
            Font,
            PatternFill,
        )

        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return ExportArtifact(
            content=output,
            filename=self._filename("xlsx"),
            mimetype=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )

    @staticmethod
    def _safe_sheet_name(text: str, used: set[str], fallback: str = "Lainnya") -> str:
        """Samarkan nama makanan jadi nama sheet Excel yang valid.

        Excel melarang : \\ / ? * [ ] dan maksimal 31 karakter.
        """
        cleaned = re.sub(r"[:\\/?*\[\]]", "-", str(text or "").strip()).strip()
        cleaned = re.sub(r"\s+", " ", cleaned).strip() or fallback
        name = cleaned[:31]
        if name not in used:
            return name
        counter = 1
        while f"{name[:29]}-{counter}" in used:
            counter += 1
        return f"{name[:29]}-{counter}"

    def xlsx_simple(
        self, since=None, scan_id=None, scan_ids=None
    ) -> ExportArtifact:
        """Export XLSX ringkas: 1 sheet per makanan + 1 sheet rekap 'Semua'.

        Tiap sheet berisi kolom: nama, makanan, confidence_percent,
        berat_gram, protein_g, lemak_g, karbohidrat_g, energi_kkal.
        Kolom ``nama`` bernomor per makanan (``bakso 1``, ``bakso 2``, ...).
        Hanya scan yang dipilih (scan_id/scan_ids) yang masuk.
        """
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill

        dataset = self.database_service.export_dataset(
            since=since, scan_id=scan_id, scan_ids=scan_ids
        )
        scans = self._order_scans(dataset["scans"], scan_ids)
        if scan_id is not None and scan_ids is None:
            scans = dataset["scans"]

        counters: dict[str, int] = {}
        by_food: dict[str, list[dict]] = {}
        for scan in scans:
            flat = self._flat_scan(scan)
            food = str(flat.get("food_name") or "-").strip() or "-"
            key = food.lower()
            counters[key] = counters.get(key, 0) + 1
            row = {
                "nama": f"{key} {counters[key]}",
                "makanan": food,
                "confidence_percent": round(
                    float(flat.get("confidence_percent") or 0), 2
                ),
                "berat_gram": round(float(flat.get("weight_grams") or 0), 2),
                "protein_g": round(float(flat.get("protein_g_total") or 0), 2),
                "lemak_g": round(float(flat.get("fat_g_total") or 0), 2),
                "karbohidrat_g": round(
                    float(flat.get("carbohydrate_g_total") or 0), 2
                ),
                "energi_kkal": round(
                    float(flat.get("energy_kcal_total") or 0), 2
                ),
            }
            by_food.setdefault(key, []).append(row)

        workbook = Workbook()
        used_names: set[str] = set()
        first = True
        # Sheet rekap berisi semua baris terpilih.
        all_rows = [row for rows in by_food.values() for row in rows]
        for key in sorted(by_food):
            rows = by_food[key]
            display = rows[0]["makanan"]
            if first:
                sheet = workbook.active
                sheet.title = self._safe_sheet_name("Semua", used_names)
                used_names.add(sheet.title)
                self._fill_sheet(sheet, all_rows, Font, PatternFill)
                first = False
            sheet = workbook.create_sheet(
                self._safe_sheet_name(display, used_names)
            )
            used_names.add(sheet.title)
            self._fill_sheet(sheet, rows, Font, PatternFill)
        if first:  # Tidak ada data: tetap tulis 1 sheet kosong.
            sheet = workbook.active
            sheet.title = self._safe_sheet_name("Semua", used_names)
            self._fill_sheet(sheet, [], Font, PatternFill)

        output = BytesIO()
        workbook.save(output)
        output.seek(0)
        return ExportArtifact(
            content=output,
            filename=self._filename("xlsx"),
            mimetype=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
        )

    @staticmethod
    def _fill_sheet(sheet, records, Font, PatternFill) -> None:
        if not records:
            sheet.append(["Tidak ada data"])
            return
        headers = list(records[0])
        sheet.append(headers)
        for record in records:
            sheet.append([record.get(header) for header in headers])
        fill = PatternFill("solid", fgColor="163A4A")
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = fill
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            values = [str(cell.value or "") for cell in column[:100]]
            width = min(45, max(10, max(map(len, values), default=10) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width

    def _thumbnail(
        self,
        relative_path: str | None,
        max_width: float,
        max_height: float,
        caption: str,
        style,
    ):
        """Bikin flowable thumbnail untuk satu foto scan.

        Falls back to teks path bila file hilang atau tidak bisa dibaca,
        supaya export tetap berhasil untuk data legacy.
        """
        from PIL import Image as PILImage
        from reportlab.platypus import Image, Paragraph, Table

        label = Paragraph(html.escape(caption), style)
        source = (
            (self.database_service.base_dir / relative_path)
            if relative_path
            else None
        )

        try:
            if source is None or not source.is_file():
                raise FileNotFoundError(str(source))
            with PILImage.open(source) as picture:
                # 3x titik reportlab memberi sekitar 216 dpi, cukup untuk PDF.
                picture.thumbnail(
                    (int(max_width * 3), int(max_height * 3)),
                    PILImage.LANCZOS,
                )
                width_px, height_px = picture.size
                buffer = BytesIO()
                picture.convert("RGB").save(buffer, format="JPEG", quality=80)
        except Exception:
            return Table(
                [
                    [label],
                    [
                        Paragraph(
                            html.escape(
                                f"{relative_path or '-'} (file not found)"
                            ),
                            style,
                        )
                    ],
                ],
                colWidths=[max_width],
            )

        image = Image(buffer)
        # Skala agar foto muat di kotak, termasuk foto berorientasi potret.
        scale = min(max_width / width_px, max_height / height_px)
        image.drawWidth = width_px * scale
        image.drawHeight = height_px * scale
        return Table([[label], [image]], colWidths=[max_width])

    def pdf(self, since=None, scan_id=None, scan_ids=None) -> ExportArtifact:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4, landscape
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.platypus import (
            PageBreak,
            Paragraph,
            SimpleDocTemplate,
            Spacer,
            Table,
            TableStyle,
        )

        dataset = self.database_service.export_dataset(
            since=since, scan_id=scan_id, scan_ids=scan_ids
        )
        dataset["scans"] = self._order_scans(dataset["scans"], scan_ids)
        detections_by_scan: dict[str, list[dict]] = {}
        for detection in dataset["detections"]:
            detections_by_scan.setdefault(detection["scan_id"], []).append(
                detection
            )

        output = BytesIO()
        document = SimpleDocTemplate(
            output,
            pagesize=landscape(A4),
            leftMargin=12 * mm,
            rightMargin=12 * mm,
            topMargin=12 * mm,
            bottomMargin=12 * mm,
            title="Food Recognition Report",
        )
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "ReportTitle",
            parent=styles["Title"],
            alignment=TA_CENTER,
            textColor=colors.HexColor("#163A4A"),
        )
        caption_style = ParagraphStyle(
            "PhotoCaption",
            parent=styles["Normal"],
            fontSize=8,
            textColor=colors.HexColor("#163A4A"),
        )
        photo_width = 88 * mm
        photo_height = 58 * mm
        story = [
            Paragraph("Food Recognition Report", title_style),
            Paragraph(
                f"Created on {datetime.now().strftime('%d-%m-%Y %H:%M:%S')} · "
                f"{len(dataset['scans'])} scan",
                styles["Normal"],
            ),
            Spacer(1, 6 * mm),
        ]

        if not dataset["scans"]:
            story.append(Paragraph("No scan data yet.", styles["Normal"]))

        for index, scan in enumerate(dataset["scans"]):
            flat = self._flat_scan(scan)
            nutrition_status = (
                "Available" if flat["nutrition_available"] else "Not available"
            )
            details = [
                ["Scan ID", flat["scan_id"], "Timestamp", flat["timestamp"]],
                ["Food", flat["food_name"], "Source", flat["source"]],
                [
                    "Confidence",
                    f"{float(flat['confidence_percent'] or 0):.2f}%",
                    "Weight",
                    f"{float(flat['weight_grams'] or 0):.2f} g",
                ],
                [
                    "Objects",
                    flat["object_count"],
                    "Model",
                    f"{flat['model_name'] or '-'} / {flat['model_input'] or '-'}",
                ],
                [
                    "Inference Time",
                    f"{float(flat['inference_ms'] or 0):.2f} ms",
                    "Total Processing Time",
                    f"{float(flat['complete_ms'] or 0):.2f} ms",
                ],
                ["Nutrition", nutrition_status, "Version", flat["nutrition_data_version"] or "-"],
                [
                    "Total Energy",
                    f"{float(flat['energy_kcal_total'] or 0):.3f} kcal",
                    "Protein / Fat / Carbs",
                    (
                        f"{float(flat['protein_g_total'] or 0):.3f} / "
                        f"{float(flat['fat_g_total'] or 0):.3f} / "
                        f"{float(flat['carbohydrate_g_total'] or 0):.3f} g"
                    ),
                ],
            ]
            table = Table(details, colWidths=[28 * mm, 78 * mm, 34 * mm, 110 * mm])
            table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#E7F2F5")),
                        ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#E7F2F5")),
                        ("FONTNAME", (0, 0), (0, -1), "Helvetica-Bold"),
                        ("FONTNAME", (2, 0), (2, -1), "Helvetica-Bold"),
                        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#9CB9C4")),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                        ("PADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )
            story.extend([table, Spacer(1, 4 * mm)])

            photos = Table(
                [
                    [
                        self._thumbnail(
                            flat["original_image_path"],
                            photo_width,
                            photo_height,
                            "Image Before Detection",
                            caption_style,
                        ),
                        self._thumbnail(
                            flat["annotated_image_path"],
                            photo_width,
                            photo_height,
                            "Detection Result Image",
                            caption_style,
                        ),
                    ]
                ],
                colWidths=[photo_width, photo_width],
            )
            photos.setStyle(
                TableStyle(
                    [
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 0),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 8),
                        ("TOPPADDING", (0, 0), (-1, -1), 0),
                    ]
                )
            )
            story.extend([photos, Spacer(1, 4 * mm)])

            detection_rows = [["Class ID", "Name", "Confidence", "Bounding Box"]]
            for detection in detections_by_scan.get(scan["scan_id"], []):
                detection_rows.append(
                    [
                        detection["class_id"],
                        detection["class_name"],
                        f"{float(detection['confidence_percent']):.2f}%",
                        (
                            f"({detection['x1']}, {detection['y1']}) - "
                            f"({detection['x2']}, {detection['y2']})"
                        ),
                    ]
                )
            if len(detection_rows) == 1:
                detection_rows.append(["-", "No detections", "-", "-"])
            detection_table = Table(
                detection_rows,
                colWidths=[25 * mm, 80 * mm, 35 * mm, 70 * mm],
            )
            detection_table.setStyle(
                TableStyle(
                    [
                        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#163A4A")),
                        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                        ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#9CB9C4")),
                        ("FONTSIZE", (0, 0), (-1, -1), 8),
                        ("PADDING", (0, 0), (-1, -1), 4),
                    ]
                )
            )
            story.append(detection_table)
            if index < len(dataset["scans"]) - 1:
                story.append(PageBreak())

        document.build(story)
        output.seek(0)
        return ExportArtifact(
            content=output,
            filename=self._filename("pdf"),
            mimetype="application/pdf",
        )

    def create(
        self, export_format: str, since=None, scan_id=None, scan_ids=None
    ) -> ExportArtifact:
        creators = {
            "csv": self.csv,
            "csv_simple": self.csv_simple,
            "csv-ringkas": self.csv_simple,
            "ringkas": self.csv_simple,
            "xlsx_simple": self.xlsx_simple,
            "xlsx-ringkas": self.xlsx_simple,
            "xlsx": self.xlsx,
            "pdf": self.pdf,
            "jpg": self.jpg,
            "jpeg": self.jpg,
        }
        try:
            creator = creators[export_format.lower()]
        except KeyError as error:
            raise ValueError(
                "Format export harus csv, xlsx, pdf, jpg, atau jpeg."
            ) from error
        return creator(since=since, scan_id=scan_id, scan_ids=scan_ids)

    @staticmethod
    def _safe_scan_filename(scan_id: str | None, index: int) -> str:
        base = re.sub(r"[^A-Za-z0-9_-]+", "_", str(scan_id or f"scan-{index}"))
        return f"{base or f'scan-{index}'}.jpg"

    def csv_simple(
        self, since=None, scan_id=None, scan_ids=None
    ) -> ExportArtifact:
        """Export CSV ringkas: hanya kolom yang diminta user.

        Kolom: nama, makanan, confidence_percent, berat_gram,
        protein_g, lemak_g, karbohidrat_g, energi_kkal.
        Kolom ``nama`` otomatis bernomor per makanan
        (``bakso 1``, ``bakso 2``, ...) sesuai urutan pilihan.
        Hanya scan yang dipilih (scan_id/scan_ids) yang masuk.
        """
        dataset = self.database_service.export_dataset(
            since=since, scan_id=scan_id, scan_ids=scan_ids
        )
        scans = self._order_scans(dataset["scans"], scan_ids)
        # Bila satu scan_id saja, pastikan urutan tetap (export_dataset
        # mengembalikan list 1 elemen, _order_scans tidak mengubahnya).
        if scan_id is not None and scan_ids is None:
            scans = dataset["scans"]

        counters: dict[str, int] = {}
        rows: list[dict] = []
        for scan in scans:
            flat = self._flat_scan(scan)
            food = str(flat.get("food_name") or "-").strip() or "-"
            key = food.lower()
            counters[key] = counters.get(key, 0) + 1
            rows.append(
                {
                    "nama": f"{key} {counters[key]}",
                    "makanan": food,
                    "confidence_percent": round(
                        float(flat.get("confidence_percent") or 0), 2
                    ),
                    "berat_gram": round(float(flat.get("weight_grams") or 0), 2),
                    "protein_g": round(
                        float(flat.get("protein_g_total") or 0), 2
                    ),
                    "lemak_g": round(float(flat.get("fat_g_total") or 0), 2),
                    "karbohidrat_g": round(
                        float(flat.get("carbohydrate_g_total") or 0), 2
                    ),
                    "energi_kkal": round(
                        float(flat.get("energy_kcal_total") or 0), 2
                    ),
                }
            )

        output = StringIO(newline="")
        fieldnames = [
            "nama",
            "makanan",
            "confidence_percent",
            "berat_gram",
            "protein_g",
            "lemak_g",
            "karbohidrat_g",
            "energi_kkal",
        ]
        writer = csv.DictWriter(output, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)
        content = BytesIO(("\ufeff" + output.getvalue()).encode("utf-8"))
        return ExportArtifact(
            content=content,
            filename=self._filename("csv"),
            mimetype="text/csv; charset=utf-8",
        )

    @staticmethod
    def _hasil_row_from_record(record: dict) -> dict:
        """Satu baris hasil.csv dari record ScanService (kolom nama diisi saat tulis)."""
        nutrition = record.get("nutrition") or {}
        totals = nutrition.get("totals") or {}
        food = str(record.get("food_name") or "-").strip() or "-"
        confidence = record.get("confidence_percent", record.get("confidence", 0))
        weight = record.get("weight_grams", record.get("weight", 0))

        def _num(value) -> float:
            try:
                return round(float(value or 0), 2)
            except (TypeError, ValueError):
                return 0.0

        return {
            "makanan": food,
            "confidence_percent": _num(confidence),
            "berat_gram": _num(weight),
            "protein_g": _num(totals.get("protein_g")),
            "lemak_g": _num(totals.get("fat_g")),
            "karbohidrat_g": _num(totals.get("carbohydrate_g")),
            "energi_kkal": _num(totals.get("energy_kcal")),
        }

    @staticmethod
    def _hasil_row_from_flat(flat: dict) -> dict:
        """Satu baris hasil.csv dari flat scan DB (kolom nama diisi saat tulis)."""

        def _num(value) -> float:
            try:
                return round(float(value or 0), 2)
            except (TypeError, ValueError):
                return 0.0

        food = str(flat.get("food_name") or "-").strip() or "-"
        return {
            "makanan": food,
            "confidence_percent": _num(flat.get("confidence_percent")),
            "berat_gram": _num(flat.get("weight_grams")),
            "protein_g": _num(flat.get("protein_g_total")),
            "lemak_g": _num(flat.get("fat_g_total")),
            "karbohidrat_g": _num(flat.get("carbohydrate_g_total")),
            "energi_kkal": _num(flat.get("energy_kcal_total")),
        }

    @staticmethod
    def _read_hasil_rows(csv_path) -> list[dict]:
        """Baca baris-baris hasil.csv secara toleran (file boleh belum ada)."""
        from pathlib import Path as _Path

        path = _Path(csv_path)
        if not path.is_file() or path.stat().st_size == 0:
            return []
        numeric_keys = (
            "confidence_percent",
            "berat_gram",
            "protein_g",
            "lemak_g",
            "karbohidrat_g",
            "energi_kkal",
        )
        rows: list[dict] = []
        try:
            with path.open("r", encoding="utf-8-sig", newline="") as handle:
                for raw in csv.DictReader(handle):
                    if not raw:
                        continue
                    food = str(raw.get("makanan") or "-").strip() or "-"
                    row: dict = {"makanan": food}
                    for key in numeric_keys:
                        try:
                            row[key] = round(float(raw.get(key) or 0), 2)
                        except (TypeError, ValueError):
                            row[key] = 0.0
                    rows.append(row)
        except (OSError, csv.Error):
            return []
        return rows

    @staticmethod
    def _write_hasil_rows(rows: list[dict], csv_path):
        """Tulis ulang hasil.csv: urut per class + penomoran nama per class.

        Contoh nama: ``bakso 1``, ``bakso 2``, ... Urutan dalam satu class
        mengikuti urutan baris masuk (stabil).
        """
        from pathlib import Path as _Path

        path = _Path(csv_path)
        path.parent.mkdir(parents=True, exist_ok=True)
        ordered = sorted(rows, key=lambda item: str(item.get("makanan") or "-").lower())
        counters: dict[str, int] = {}
        final: list[dict] = []
        for row in ordered:
            food = str(row.get("makanan") or "-").strip() or "-"
            key = food.lower()
            counters[key] = counters.get(key, 0) + 1
            final.append(
                {
                    "nama": f"{key} {counters[key]}",
                    "makanan": food,
                    "confidence_percent": row.get("confidence_percent", 0.0),
                    "berat_gram": row.get("berat_gram", 0.0),
                    "protein_g": row.get("protein_g", 0.0),
                    "lemak_g": row.get("lemak_g", 0.0),
                    "karbohidrat_g": row.get("karbohidrat_g", 0.0),
                    "energi_kkal": row.get("energi_kkal", 0.0),
                }
            )
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=ExportService.SINGLE_CSV_FIELDNAMES)
            writer.writeheader()
            writer.writerows(final)
        return path

    @staticmethod
    def _migrate_legacy_csv(legacy_path) -> list[dict]:
        """Pindahkan isi file lama (hasil_scan.csv) lalu hapus file lama.

        Return baris-baris lama agar tidak ada data yang hilang.
        """
        from pathlib import Path as _Path

        if not legacy_path:
            return []
        legacy = _Path(legacy_path)
        if not legacy.is_file():
            return []
        moved = ExportService._read_hasil_rows(legacy)
        try:
            legacy.unlink()
        except OSError:
            pass
        return moved

    def update_hasil_csv(self, records, csv_path, legacy_path=None):
        """Gabung records baru ke hasil.csv lalu tulis ulang urut per class.

        Dipakai tiap scan sukses. Isi file lama (legacy_path) dimigrasi
        dulu bila masih ada agar tak ada data yang hilang.
        """
        if isinstance(records, dict):
            items = [records]
        else:
            items = list(records or [])
        new_rows = [self._hasil_row_from_record(record) for record in items]
        legacy_rows = self._migrate_legacy_csv(legacy_path)
        all_rows = self._read_hasil_rows(csv_path) + legacy_rows + new_rows
        return self._write_hasil_rows(all_rows, csv_path)

    def rebuild_hasil_from_scans(self, scans, csv_path, legacy_path=None):
        """Tulis ulang hasil.csv dari daftar scan database, urut per class."""
        rows = [self._hasil_row_from_flat(self._flat_scan(scan)) for scan in (scans or [])]
        legacy_rows = self._migrate_legacy_csv(legacy_path)
        return self._write_hasil_rows(legacy_rows + rows, csv_path)

    def _load_export_photo_by_key(self, flat: dict, key: str):
        """Ambil satu foto (original/annotated) sebagai PIL Image (RGB) atau None."""
        from PIL import Image as PILImage

        relative_path = flat.get(key)
        if not relative_path:
            return None
        source = self.database_service.base_dir / relative_path
        try:
            if not source.is_file():
                return None
            with PILImage.open(source) as picture:
                return picture.convert("RGB").copy()
        except Exception:
            return None

    def _load_export_photo(self, flat: dict):
        """Ambil foto annotated/original sebagai PIL Image (RGB) atau None."""
        photo = self._load_export_photo_by_key(flat, "annotated_image_path")
        if photo is not None:
            return photo
        return self._load_export_photo_by_key(flat, "original_image_path")

    @staticmethod
    def _wrap_text(draw, text: str, font, max_width: int) -> list[str]:
        """Bungkus teks per kata agar muat di max_width (PIL)."""
        words = str(text).split()
        if not words:
            return [""]
        lines: list[str] = []
        current = ""
        for word in words:
            candidate = f"{current} {word}".strip()
            try:
                width = draw.textlength(candidate, font=font)
            except Exception:
                width = len(candidate) * 12
            if width <= max_width or not current:
                current = candidate
            else:
                lines.append(current)
                current = word
        if current:
            lines.append(current)
        return lines

    def _render_scan_card(self, scan: dict, detections: list[dict]) -> BytesIO:
        """Render 1 kartu laporan JPG dengan isi yang sama seperti PDF.

        Isi per scan (mirip tabel PDF):
        - Header: nama makanan + timestamp + scan ID.
        - Dua foto berdampingan: sebelum deteksi (original) & hasil deteksi
          (annotated), sama seperti PDF.
        - 7 baris detail: scan ID/waktu, makanan/sumber, confidence/berat,
          objek/model, inferensi/total, nutrisi/versi, energi/P-L-K.
        - Daftar deteksi lengkap: nama, confidence, bounding box
          (bukan hanya top-2).
        """
        from PIL import Image as PILImage, ImageDraw, ImageFont

        flat = self._flat_scan(scan)
        photo_original = self._load_export_photo_by_key(flat, "original_image_path")
        photo_annotated = self._load_export_photo_by_key(flat, "annotated_image_path")

        try:
            font_title = ImageFont.load_default(size=46)
            font_sub = ImageFont.load_default(size=27)
            font_body = ImageFont.load_default(size=30)
            font_small = ImageFont.load_default(size=25)
            font_caption = ImageFont.load_default(size=26)
        except TypeError:  # Pillow lama tanpa argumen size.
            font_title = ImageFont.load_default()
            font_sub = font_title
            font_body = font_title
            font_small = font_title
            font_caption = font_title

        width = 1600
        pad = 56
        header_h = 175
        photo_h = 620
        caption_h = 46
        col_gap = 32
        line_h = 42
        row_gap = 14
        section_gap = 28
        footer_h = 70

        col_width = (width - pad * 2 - col_gap) // 2
        full_width = width - pad * 2

        nutrition_status = (
            "Available" if flat.get("nutrition_available") else "Not available"
        )
        detail_rows = [
            (
                f"Scan ID: {flat.get('scan_id') or '-'}",
                f"Timestamp: {flat.get('timestamp') or '-'}",
            ),
            (
                f"Food: {flat.get('food_name') or '-'}",
                f"Source: {flat.get('source') or '-'}",
            ),
            (
                f"Confidence: {float(flat.get('confidence_percent') or 0):.2f}%",
                f"Weight: {float(flat.get('weight_grams') or 0):.2f} g",
            ),
            (
                f"Objects: {flat.get('object_count') or 0}",
                f"Model: {flat.get('model_name') or '-'} / "
                f"{flat.get('model_input') or '-'}",
            ),
            (
                f"Inference Time: {float(flat.get('inference_ms') or 0):.2f} ms",
                f"Total Processing Time: {float(flat.get('complete_ms') or 0):.2f} ms",
            ),
            (
                f"Nutrition: {nutrition_status}",
                f"Version: {flat.get('nutrition_data_version') or '-'}",
            ),
            (
                f"Total Energy: {float(flat.get('energy_kcal_total') or 0):.3f} kcal",
                f"Protein / Fat / Carbs: "
                f"{float(flat.get('protein_g_total') or 0):.3f} / "
                f"{float(flat.get('fat_g_total') or 0):.3f} / "
                f"{float(flat.get('carbohydrate_g_total') or 0):.3f} g",
            ),
        ]

        if detections:
            detection_lines = [
                f"- {item.get('class_name', '-')} "
                f"({float(item.get('confidence_percent', 0)):.2f}%) "
                f"({item.get('x1')}, {item.get('y1')}) - "
                f"({item.get('x2')}, {item.get('y2')})"
                for item in sorted(
                    detections,
                    key=lambda item: float(item.get("confidence_percent", 0)),
                    reverse=True,
                )
            ]
        else:
            detection_lines = ["- No detections"]

        # Hitung tinggi dinamis dengan wrapping terlebih dahulu.
        measurer = ImageDraw.Draw(PILImage.new("RGB", (width, 10), "white"))
        wrapped_rows: list[tuple[list[str], list[str]]] = []
        for left, right in detail_rows:
            wrapped_rows.append(
                (
                    self._wrap_text(measurer, left, font_body, col_width),
                    self._wrap_text(measurer, right, font_body, col_width),
                )
            )
        wrapped_detections: list[list[str]] = [
            self._wrap_text(measurer, line, font_body, full_width)
            for line in detection_lines
        ]

        details_h = sum(
            max(len(left), len(right)) * line_h + row_gap
            for left, right in wrapped_rows
        )
        detection_h = (
            line_h  # title "Detections (n)"
            + sum(len(lines) * line_h for lines in wrapped_detections)
            + row_gap
        )
        height = (
            header_h
            + caption_h
            + photo_h
            + section_gap
            + details_h
            + section_gap
            + detection_h
            + section_gap
            + footer_h
        )

        card = PILImage.new("RGB", (width, height), "white")
        draw = ImageDraw.Draw(card)

        # Header.
        draw.rectangle([0, 0, width, header_h], fill="#163A4A")
        draw.text(
            (pad, 24),
            "Food Recognition Report",
            font=font_sub,
            fill="#C7DEE7",
        )
        draw.text(
            (pad, 62),
            str(flat.get("food_name") or "Food Recognition"),
            font=font_title,
            fill="white",
        )
        draw.text(
            (pad, header_h - 42),
            f"{flat.get('timestamp') or '-'}  ·  #{flat.get('scan_id') or '-'}",
            font=font_sub,
            fill="#C7DEE7",
        )

        # Dua foto berdampingan (sama urutannya dengan PDF).
        y_photos = header_h
        for col_index, (caption, photo) in enumerate(
            [
                ("Image Before Detection", photo_original),
                ("Detection Result Image", photo_annotated),
            ]
        ):
            x0 = pad + col_index * (col_width + col_gap)
            draw.text((x0, y_photos + 8), caption, font=font_caption, fill="#163A4A")
            box = PILImage.new("RGB", (col_width, photo_h), "#0E222B")
            if photo is not None:
                fitted = photo.copy()
                fitted.thumbnail((col_width, photo_h), PILImage.LANCZOS)
                box.paste(
                    fitted,
                    (
                        (col_width - fitted.size[0]) // 2,
                        (photo_h - fitted.size[1]) // 2,
                    ),
                )
            else:
                empty = ImageDraw.Draw(box)
                text = "Image not found"
                try:
                    text_w = empty.textlength(text, font=font_body)
                except Exception:
                    text_w = len(text) * 12
                empty.text(
                    ((col_width - text_w) / 2, (photo_h - 30) / 2),
                    text,
                    font=font_body,
                    fill="white",
                )
            card.paste(box, (x0, y_photos + caption_h))

        # Detail (2 kolom per baris, sama seperti tabel PDF).
        y = y_photos + caption_h + photo_h + section_gap
        for left_lines, right_lines in wrapped_rows:
            rows_in_line = max(len(left_lines), len(right_lines))
            for i in range(rows_in_line):
                if i < len(left_lines):
                    draw.text((pad, y), left_lines[i], font=font_body, fill="#163A4A")
                if i < len(right_lines):
                    draw.text(
                        (pad + col_width + col_gap, y),
                        right_lines[i],
                        font=font_body,
                        fill="#163A4A",
                    )
                y += line_h
            y += row_gap

        # Full detections.
        y += section_gap - row_gap
        draw.text(
            (pad, y),
            f"Detections ({len(detections)}):",
            font=font_body,
            fill="#163A4A",
        )
        y += line_h
        for lines in wrapped_detections:
            for line in lines:
                draw.text((pad, y), line, font=font_body, fill="#163A4A")
                y += line_h
        y += row_gap

        draw.text(
            (pad, height - 48),
            f"FoodVision Analytics · {datetime.now().strftime('%d-%m-%Y %H:%M')}",
            font=font_small,
            fill="#6B8A96",
        )

        output = BytesIO()
        card.save(output, format="JPEG", quality=92)
        output.seek(0)
        return output

    def jpg(self, since=None, scan_id=None, scan_ids=None) -> ExportArtifact:
        """Export JPG.

        1 scan -> 1 file .jpg. Banyak scan -> 1 file .zip berisi kartu JPG
        per scan (mengikuti filter range yang sama seperti PDF).
        """
        dataset = self.database_service.export_dataset(
            since=since, scan_id=scan_id, scan_ids=scan_ids
        )
        dataset["scans"] = self._order_scans(dataset["scans"], scan_ids)
        detections_by_scan: dict[str, list[dict]] = {}
        for detection in dataset["detections"]:
            detections_by_scan.setdefault(detection["scan_id"], []).append(
                detection
            )

        scans = dataset["scans"]
        if len(scans) <= 1:
            if not scans:
                card = self._render_scan_card({}, [])
                card.seek(0)
                return ExportArtifact(
                    content=card,
                    filename=self._filename("jpg"),
                    mimetype="image/jpeg",
                )
            scan = scans[0]
            card = self._render_scan_card(
                scan, detections_by_scan.get(scan["scan_id"], [])
            )
            filename = self._safe_scan_filename(scan.get("scan_id"), 0)
            return ExportArtifact(
                content=card,
                filename=filename,
                mimetype="image/jpeg",
            )

        # Banyak scan: bungkus semua kartu JPG ke dalam satu ZIP.
        output = BytesIO()
        with zipfile.ZipFile(
            output, "w", compression=zipfile.ZIP_DEFLATED
        ) as archive:
            for index, scan in enumerate(scans):
                card = self._render_scan_card(
                    scan, detections_by_scan.get(scan["scan_id"], [])
                )
                archive.writestr(
                    self._safe_scan_filename(scan.get("scan_id"), index),
                    card.getvalue(),
                )
        output.seek(0)
        timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        return ExportArtifact(
            content=output,
            filename=f"food-recognition-{timestamp}.zip",
            mimetype="application/zip",
        )
